"""Build a PPTX where every slide is a full-bleed looping video that starts automatically."""
import html
import json
import re
import subprocess
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.util import Emu

W = Path(sys.argv[1])
OUT = Path(sys.argv[2])
deck = json.loads((W / "live/project/deck.json").read_text(encoding="utf-8"))
loops = dict(l.split() for l in (W / "jobs.txt").read_text().split("\n") if l.strip())

P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"

TIMING = """<p:timing xmlns:p="{ns}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>
<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="5" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="afterEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="6" dur="{dur}" fill="hold"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:cmd>
</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>
</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>
<p:video><p:cMediaNode vol="0" mute="1"><p:cTn id="7" repeatCount="indefinite" fill="hold" display="0"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cMediaNode></p:video>
</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>"""

prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
blank = prs.slide_layouts[6]
vid_dir = W / "video"
vid_dir.mkdir(exist_ok=True)

for sid in deck["order"]:
    frames = W / "frames" / sid
    mp4 = vid_dir / f"{sid}.mp4"
    subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-framerate", "25", "-i", str(frames / "f%04d.jpg"),
                    "-c:v", "libx264", "-preset", "slow", "-crf", "20", "-pix_fmt", "yuv420p",
                    "-movflags", "+faststart", "-an", str(mp4)], check=True)
    poster = frames / "f0000.jpg"
    slide = prs.slides.add_slide(blank)
    mv = slide.shapes.add_movie(str(mp4), 0, 0, prs.slide_width, prs.slide_height,
                                poster_frame_image=str(poster), mime_type="video/mp4")
    sld = slide._element
    for t in sld.findall(f"{{{P_NS}}}timing"):
        sld.remove(t)
    dur_ms = int(float(loops[sid]) * 1000)
    timing = etree.fromstring(TIMING.format(ns=P_NS, dur=dur_ms, spid=mv.shape_id))
    # p:timing must come after p:clrMapOvr / before p:extLst
    ext = sld.find(f"{{{P_NS}}}extLst")
    if ext is not None:
        ext.addprevious(timing)
    else:
        sld.append(timing)
    src = (W / f"live/project/slides/{sid}.html").read_text(encoding="utf-8")
    m = re.search(r"<aside>(.*?)</aside>", src, re.S)
    if m:
        slide.notes_slide.notes_text_frame.text = html.unescape(m.group(1).strip())
    print(sid, mp4.stat().st_size // 1024, "KB")

prs.save(OUT)
print("saved", OUT, OUT.stat().st_size // 1024, "KB")
