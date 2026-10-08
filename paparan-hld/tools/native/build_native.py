"""Assemble a PPTX from extract.js output using native PowerPoint shapes, text and built-in animation timing."""
import json
import math
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu

OUTDIR = Path(sys.argv[1])          # native/out
ORDER = json.loads(Path(sys.argv[2]).read_text())["order"]
NOTES = json.loads(Path(sys.argv[3]).read_text(encoding="utf-8"))
DEST = Path(sys.argv[4])
CAL = json.loads(Path(sys.argv[5]).read_text()) if len(sys.argv) > 5 and Path(sys.argv[5]).exists() else {}

PX = 6350  # EMU per CSS px (1920px == 12192000 EMU)
SW, SH = 1920, 1080
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def E(v):
    return Emu(int(round(v * PX)))


def hexc(c):
    return "%02X%02X%02X" % tuple(int(round(x)) for x in c[:3])


# ---------------- keyframe utilities ----------------
def rdp(points, eps, dims):
    """Ramer-Douglas-Peucker over (t, v1..vn) keeping the time column."""
    if len(points) < 3:
        return points
    a, b = points[0], points[-1]
    dmax, idx = 0, 0
    for i in range(1, len(points) - 1):
        p = points[i]
        f = (p[0] - a[0]) / ((b[0] - a[0]) or 1)
        d = max(abs(p[k] - (a[k] + (b[k] - a[k]) * f)) / dims[k - 1] for k in range(1, len(p)))
        if d > dmax:
            dmax, idx = d, i
    if dmax > eps:
        return rdp(points[: idx + 1], eps, dims)[:-1] + rdp(points[idx:], eps, dims)
    return [a, b]


class Timing:
    def __init__(self):
        self.next_id = 5
        self.effects = []

    def nid(self):
        self.next_id += 1
        return self.next_id

    def cbhvr(self, spid, dur, delay=0, attr=None, additive=None, fill="hold"):
        st = f'<p:stCondLst><p:cond delay="{int(delay)}"/></p:stCondLst>' if delay else ""
        add = f' additive="{additive}"' if additive else ""
        an = f"<p:attrNameLst><p:attrName>{attr}</p:attrName></p:attrNameLst>" if attr else ""
        return (f'<p:cBhvr{add}><p:cTn id="{self.nid()}" dur="{int(max(1, dur))}" fill="{fill}">{st}</p:cTn>'
                f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>{an}</p:cBhvr>')

    def anim(self, spid, attr, keys, L):
        tav = "".join(f'<p:tav tm="{int(round(t / L * 100000))}"><p:val><p:strVal val="{v:.6f}"/></p:val></p:tav>' for t, v in keys)
        return (f'<p:anim calcmode="lin" valueType="num">{self.cbhvr(spid, L * 1000, attr=attr, additive="base")}'
                f"<p:tavLst>{tav}</p:tavLst></p:anim>")

    def rot(self, spid, by_deg, dur_ms, delay_ms):
        return f'<p:animRot by="{int(round(by_deg * 60000))}">{self.cbhvr(spid, dur_ms, delay_ms, attr="r")}</p:animRot>'

    def setvis(self, spid, val, delay_ms):
        return (f'<p:set>{self.cbhvr(spid, 1, delay_ms, attr="style.visibility")}'
                f'<p:to><p:strVal val="{val}"/></p:to></p:set>')

    def fade(self, spid, inout, dur_ms, delay_ms):
        return f'<p:animEffect transition="{inout}" filter="fade">{self.cbhvr(spid, dur_ms, delay_ms, fill="hold")}</p:animEffect>'

    def effect(self, behaviors, period_ms):
        hidden0 = any('style.visibility' in b and 'val="hidden"' in b and '<p:cond delay=' not in b for b in behaviors)
        cls = 'presetID="1" presetClass="entr" presetSubtype="0"' if hidden0 else 'presetID="0" presetClass="emph" presetSubtype="0"'
        self.effects.append(
            f'<p:par><p:cTn id="{self.nid()}" {cls} dur="{int(period_ms)}" repeatCount="indefinite" fill="hold" grpId="0" nodeType="withEffect">'
            f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>{"".join(behaviors)}</p:childTnLst></p:cTn></p:par>')

    def xml(self):
        if not self.effects:
            return None
        return (f'<p:timing xmlns:p="{P_NS}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
                '<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
                '<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>'
                '<p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
                + "".join(self.effects) +
                '</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>'
                '</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
                '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
                '</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>')


def opacity_behaviors(tm, spid, ts, f, L):
    """Approximate an opacity curve f(t) in [0,1] with set-visibility + fade in/out effects."""
    out = []
    n = len(f)
    if min(f) > 0.5 and max(f) - min(f) < 0.15:
        return out
    vis = f[0] > 0.5
    out.append(tm.setvis(spid, "visible" if vis else "hidden", 0))
    dt = ts[1] - ts[0]
    i = 1
    while i < n:
        if (f[i] > 0.5) != (f[i - 1] > 0.5):
            rising = f[i] > 0.5
            # extend backward to plateau start
            s = i - 1
            while s > 0 and ((f[s - 1] < f[s]) if rising else (f[s - 1] > f[s])) and (f[s] > 0.02 if rising else f[s] < 0.98):
                s -= 1
            e = i
            while e < n - 1 and ((f[e + 1] > f[e]) if rising else (f[e + 1] < f[e])) and (f[e] < 0.98 if rising else f[e] > 0.02):
                e += 1
            t0, t1 = ts[s], ts[e]
            dur = max(t1 - t0, 0)
            if dur <= dt * 1.01:
                out.append(tm.setvis(spid, "visible" if rising else "hidden", t0 * 1000 + (dur * 1000 if not rising else 0)))
            elif rising:
                out.append(tm.setvis(spid, "visible", t0 * 1000))
                out.append(tm.fade(spid, "in", dur * 1000, t0 * 1000))
            else:
                out.append(tm.fade(spid, "out", dur * 1000, t0 * 1000))
                out.append(tm.setvis(spid, "hidden", t1 * 1000 - 1))
            i = e + 1
        else:
            i += 1
    return out


def motion_behaviors(tm, spid, ts, X, Y, Wd, Ht, R, L):
    out = []
    pts = [(t, x, y, w, h) for t, x, y, w, h in zip(ts, X, Y, Wd, Ht)]
    keep = rdp(pts, 0.6, [1, 1, 1, 1])
    if max(X) - min(X) > 0.5:
        out.append(tm.anim(spid, "ppt_x", [(p[0], p[1] / SW) for p in keep], L))
    if max(Y) - min(Y) > 0.5:
        out.append(tm.anim(spid, "ppt_y", [(p[0], p[2] / SH) for p in keep], L))
    if max(Wd) - min(Wd) > 0.5:
        out.append(tm.anim(spid, "ppt_w", [(p[0], max(p[3], 0.05) / SW) for p in keep], L))
    if max(Ht) - min(Ht) > 0.5:
        out.append(tm.anim(spid, "ppt_h", [(p[0], max(p[4], 0.05) / SH) for p in keep], L))
    if R and max(R) - min(R) > 0.5:
        rk = rdp([(t, r) for t, r in zip(ts, R)], 0.5, [1])
        for (t0, r0), (t1, r1) in zip(rk, rk[1:]):
            if abs(r1 - r0) > 0.01:
                out.append(tm.rot(spid, r1 - r0, (t1 - t0) * 1000, t0 * 1000))
    return out


def unwrap(a):
    out = [a[0]]
    for v in a[1:]:
        d = v - out[-1]
        d = (d + 180) % 360 - 180
        out.append(out[-1] + d)
    return out


def color_weights(samples, states):
    cols = [[int(h[k:k + 2], 16) for k in (1, 3, 5)] for h in states]
    W = []
    for s in samples:
        f = s["fill"] if isinstance(s["fill"], list) else None
        if not f:
            W.append([1.0 / len(cols)] * len(cols))
            continue
        d = [math.dist(f, c) for c in cols]
        inv = [1 / (x + 1e-3) ** 2 for x in d]
        tot = sum(inv)
        W.append([v / tot for v in inv])
    return W


# ---------------- text ----------------
FONTMAP = {
    ("IBM Plex Sans", 400): ("IBM Plex Sans", False), ("IBM Plex Sans", 500): ("IBM Plex Sans Medium", False),
    ("IBM Plex Sans", 600): ("IBM Plex Sans SemiBold", False), ("IBM Plex Sans", 700): ("IBM Plex Sans", True),
    ("Space Grotesk", 400): ("Space Grotesk", False), ("Space Grotesk", 500): ("Space Grotesk Medium", False),
    ("Space Grotesk", 600): ("Space Grotesk SemiBold", False), ("Space Grotesk", 700): ("Space Grotesk", True),
}


def font_for(tok):
    fam = tok["font"] if tok["font"] in ("IBM Plex Sans", "Space Grotesk") else "IBM Plex Sans"
    w = min((400, 500, 600, 700), key=lambda x: abs(x - tok["weight"]))
    return FONTMAP[(fam, w)]


def add_text(slide, blk):
    toks = blk["toks"]
    lines = []
    for t in toks:
        if lines and abs(t["top"] - lines[-1][0]["top"]) < t["size"] * 0.45 and t["l"] >= lines[-1][-1]["l"] - 1:
            lines[-1].append(t)
        else:
            lines.append([t])
    lh = blk["lh"]
    first = lines[0][0]
    size0 = max(t["size"] for t in lines[0])
    fam0 = first["font"]
    cal = CAL.get(fam0, {"k": 0.0})
    # CSS line-box top of the first line
    cont_h = max(t["bot"] - t["top"] for t in lines[0])
    linebox_top = min(t["top"] for t in lines[0]) - (lh - cont_h) / 2
    top = linebox_top + cal["k"] * size0
    al = blk["align"]
    minl = min(t["l"] for t in toks)
    maxr = max(t["r"] for t in toks)
    if al in ("center",):
        left, right = blk["box"]["l"], blk["box"]["r"]
        algn = PP_ALIGN.CENTER
    elif al in ("right", "end"):
        left, right = blk["box"]["l"], blk["box"]["r"]
        algn = PP_ALIGN.RIGHT
    else:
        left, right = minl, maxr + max(40, (maxr - minl) * 0.15)
        algn = PP_ALIGN.LEFT
    height = lh * len(lines)
    tb = slide.shapes.add_textbox(E(left), E(top), E(max(right - left, 4)), E(height))
    tf = tb.text_frame
    tf.word_wrap = False
    tf.auto_size = None
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, 0)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for li, line in enumerate(lines):
        p = tf.paragraphs[0] if li == 0 else tf.add_paragraph()
        p.alignment = algn
        pPr = p._p.get_or_add_pPr()
        ln = etree.SubElement(pPr, qn("a:lnSpc"))
        etree.SubElement(ln, qn("a:spcPts")).set("val", str(int(round(lh * 50))))
        sb = etree.SubElement(pPr, qn("a:spcBef"))
        etree.SubElement(sb, qn("a:spcPts")).set("val", "0")
        # merge tokens into runs
        runs = []
        prev = None
        for t in line:
            sep = ""
            if prev is not None and (t["spaceBefore"] or t["l"] - prev["r"] > t["size"] * 0.18):
                sep = " "
            key = (t["font"], t["size"], t["weight"], t["italic"], tuple(t["color"][:3]), t["ls"], (t["color"] + [1])[3])
            if runs and runs[-1][0] == key:
                runs[-1][1] += sep + t["t"]
            else:
                if sep and runs:
                    runs[-1][1] += sep
                runs.append([key, t["t"], t])
            prev = t
        for key, text, t in runs:
            r = p.add_run()
            r.text = text
            face, bold = font_for(t)
            f = r.font
            f.size = Emu(int(t["size"] * PX))  # placeholder, replaced below
            rPr = r._r.get_or_add_rPr()
            rPr.set("sz", str(int(round(t["size"] * 50))))
            rPr.set("b", "1" if bold else "0")
            if t["italic"]:
                rPr.set("i", "1")
            if t["ls"]:
                rPr.set("spc", str(int(round(t["ls"] * 50))))
            f.color.rgb = RGBColor.from_string(hexc(t["color"]))
            alpha = t["color"][3] if len(t["color"]) > 3 else 1
            if alpha < 0.999:
                sf = rPr.find(qn("a:solidFill"))
                clr = sf[0]
                etree.SubElement(clr, qn("a:alpha")).set("val", str(int(alpha * 100000)))
            for tag in ("a:latin", "a:ea", "a:cs"):
                el = rPr.find(qn(tag))
                if el is None:
                    el = etree.SubElement(rPr, qn(tag))
                el.set("typeface", face)
    return tb


# ---------------- build ----------------
prs = Presentation()
prs.slide_width, prs.slide_height = Emu(SW * PX), Emu(SH * PX)
blank = prs.slide_layouts[6]

for sid in ORDER:
    d = OUTDIR / sid
    meta = json.loads((d / "meta.json").read_text())
    L = meta["L"]
    FPS = meta["FPS"]
    slide = prs.slides.add_slide(blank)
    tm = Timing()

    def full(img):
        return slide.shapes.add_picture(str(d / img), 0, 0, prs.slide_width, prs.slide_height)

    full("bg0.jpg")

    def emit_units(back):
        for u in meta["units"]:
            if bool(u.get("back")) != back:
                continue
            if u["kind"] == "dash":
                period = u["dash"]["dur"]
                for g in u["geo"]:
                    disp = -u["dash"]["delta"] * g["scale"]
                    dash = g["dash"]
                    if "circle" in g:
                        c = g["circle"]
                        shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, E(c["cx"] - c["rx"]), E(c["cy"] - c["ry"]), E(2 * c["rx"]), E(2 * c["ry"]))
                        if g.get("fill"):
                            shp.fill.solid(); shp.fill.fore_color.rgb = RGBColor.from_string(hexc(g["fill"]))
                        else:
                            shp.fill.background()
                        style_line(shp.line, g)
                        shp.shadow.inherit = False
                        by = math.degrees(disp / max(c["rx"], 1))
                        tm.effect([tm.rot(shp.shape_id, by, period * 1000, 0)], period * 1000)
                        continue
                    for (x1, y1), (x2, y2) in g["segs"]:
                        ln = math.hypot(x2 - x1, y2 - y1)
                        if ln < 0.5:
                            continue
                        con = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2))
                        style_line(con.line, g)
                        ux, uy = (x2 - x1) / ln, (y2 - y1) / ln
                        cx, cy = (x1 + x2) / 2, (y1 + y2) / 2
                        beh = []
                        if abs(ux * disp) > 0.05:
                            beh.append(tm.anim(con.shape_id, "ppt_x", [(0, cx / SW), (period, (cx + ux * disp) / SW)], period))
                        if abs(uy * disp) > 0.05:
                            beh.append(tm.anim(con.shape_id, "ppt_y", [(0, cy / SH), (period, (cy + uy * disp) / SH)], period))
                        if beh:
                            tm.effect(beh, period * 1000)
                continue
            S = u["samples"]
            ts = [i / FPS for i in range(len(S))]
            if u["kind"] == "ellipse":
                ref = max(range(len(S)), key=lambda i: S[i]["eff"] + S[i]["w"] * 1e-6)
                s0 = S[ref]
                w0, h0 = s0["w"] * s0["s"], s0["h"] * s0["s"]
                shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, E(s0["x"] - w0 / 2), E(s0["y"] - h0 / 2), E(w0), E(h0))
                shp.shadow.inherit = False
                base_o = s0["o"]
                if s0["fill"] and isinstance(s0["fill"], list) and s0["fo"] > 0:
                    shp.fill.solid(); shp.fill.fore_color.rgb = RGBColor.from_string(hexc(s0["fill"]))
                    set_alpha(shp.fill, s0["fo"] * base_o)
                else:
                    shp.fill.background()
                if s0["stroke"] and s0["sw"] > 0:
                    shp.line.color.rgb = RGBColor.from_string(hexc(s0["stroke"]))
                    shp.line.width = E(s0["sw"] * s0["s"])
                    set_line_alpha(shp.line, s0["so"] * base_o)
                else:
                    shp.line.fill.background()
                beh = motion_behaviors(tm, shp.shape_id, ts, [s["x"] for s in S], [s["y"] for s in S],
                                       [s["w"] * s["s"] for s in S], [s["h"] * s["s"] for s in S], None, L)
                f = [min(1, s["eff"] / max(s0["eff"], 1e-6)) for s in S]
                beh += opacity_behaviors(tm, shp.shape_id, ts, f, L)
                if beh:
                    tm.effect(beh, L * 1000)
                continue
            # sprite
            W = color_weights(S, [p["state"] for p in u["pics"]]) if u["pics"][0]["state"] else None
            for pi, pic in enumerate(u["pics"]):
                r = S[pic["ref"]]
                bx = pic["box"]
                cx0, cy0 = bx["l"] + bx["w"] / 2, bx["t"] + bx["h"] / 2
                offx, offy = cx0 - r["x"], cy0 - r["y"]
                shp = slide.shapes.add_picture(str(d / pic["file"]), E(bx["l"]), E(bx["t"]), E(bx["w"]), E(bx["h"]))
                refscale_w = max(r["w"] * r["s"], 1e-6)
                refscale_h = max(r["h"] * r["s"], 1e-6)
                rot = unwrap([s["rot"] for s in S])
                X, Y, Wd, Ht, R = [], [], [], [], []
                for s, rr in zip(S, rot):
                    kx = (s["w"] * s["s"]) / refscale_w if r["w"] > 0.01 else s["s"] / r["s"]
                    ky = (s["h"] * s["s"]) / refscale_h if r["h"] > 0.01 else s["s"] / r["s"]
                    X.append(s["x"] + offx * kx)
                    Y.append(s["y"] + offy * ky)
                    Wd.append(bx["w"] * kx)
                    Ht.append(bx["h"] * ky)
                    R.append(rr - rot[pic["ref"]])
                beh = motion_behaviors(tm, shp.shape_id, ts, X, Y, Wd, Ht, R, L)
                ref_eff = max(r["eff"], 1e-6)
                f = [min(1, s["eff"] / ref_eff) for s in S]
                if W:
                    f = [ff * w[pi] for ff, w in zip(f, W)]
                beh += opacity_behaviors(tm, shp.shape_id, ts, f, L)
                if beh:
                    tm.effect(beh, L * 1000)

    def style_line(line, g):
        line.color.rgb = RGBColor.from_string(hexc(g["stroke"]))
        line.width = E(g["width"])
        ln = line._get_or_add_ln()
        if g.get("cap") == "round":
            ln.set("cap", "rnd")
        set_line_alpha(line, g["opacity"])
        if g["dash"]:
            ds = g["dash"] if len(g["dash"]) % 2 == 0 else g["dash"] * 2
            for old in ln.findall(qn("a:prstDash")):
                ln.remove(old)
            cd = etree.SubElement(ln, qn("a:custDash"))
            wpx = max(g["width"], 0.1)
            for k in range(0, len(ds), 2):
                el = etree.SubElement(cd, qn("a:ds"))
                el.set("d", str(int(ds[k] / wpx * 100000)))
                el.set("sp", str(int(ds[k + 1] / wpx * 100000)))
            # custDash must precede round/join/headEnd... move to right place
            ln.remove(cd)
            sf = ln.find(qn("a:solidFill"))
            sf.addnext(cd)

    def set_alpha(fill, a):
        if a >= 0.999:
            return
        clr = fill._xPr.find(qn("a:solidFill"))[0]
        etree.SubElement(clr, qn("a:alpha")).set("val", str(int(max(0, a) * 100000)))

    def set_line_alpha(line, a):
        if a >= 0.999:
            return
        ln = line._get_or_add_ln()
        sf = ln.find(qn("a:solidFill"))
        if sf is not None:
            etree.SubElement(sf[0], qn("a:alpha")).set("val", str(int(max(0, a) * 100000)))

    emit_units(True)
    full("bg1.png")
    emit_units(False)
    if meta.get("top"):
        full("top.png")
    for blk in meta["texts"]:
        add_text(slide, blk)

    x = tm.xml()
    if x:
        sld = slide._element
        timing = etree.fromstring(x)
        ext = sld.find(f"{{{P_NS}}}extLst")
        if ext is not None:
            ext.addprevious(timing)
        else:
            sld.append(timing)
    if sid in NOTES:
        slide.notes_slide.notes_text_frame.text = NOTES[sid]
    print(sid, "effects", len(tm.effects), "shapes", len(slide.shapes))

prs.save(DEST)
print("saved", DEST, DEST.stat().st_size // 1024, "KB")


# ---------------- embed fonts (OFL, installable) ----------------
def embed_fonts(path, fontdir):
    import shutil, zipfile, re as _re
    fams = {
        "IBM Plex Sans": {"regular": "IBMPlexSans-Regular.ttf", "bold": "IBMPlexSansBold-Bold.ttf"},
        "IBM Plex Sans Medium": {"regular": "IBMPlexSansMedium-Regular.ttf"},
        "IBM Plex Sans SemiBold": {"regular": "IBMPlexSansSemiBold-Regular.ttf"},
        "Space Grotesk": {"regular": "SpaceGrotesk-Regular.ttf", "bold": "SpaceGroteskBold-Bold.ttf"},
        "Space Grotesk Medium": {"regular": "SpaceGroteskMedium-Regular.ttf"},
        "Space Grotesk SemiBold": {"regular": "SpaceGroteskSemiBold-Regular.ttf"},
    }
    tmp = path.with_suffix(".tmp.pptx")
    zin = zipfile.ZipFile(path)
    zout = zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED)
    rels = zin.read("ppt/_rels/presentation.xml.rels").decode()
    pres = zin.read("ppt/presentation.xml").decode()
    ct = zin.read("[Content_Types].xml").decode()
    entries, n = [], 0
    for fam, styles in fams.items():
        parts = []
        for style, fn in styles.items():
            n += 1
            rid = f"rIdFont{n}"
            zout.write(Path(fontdir) / fn, f"ppt/fonts/font{n}.fntdata")
            rels = rels.replace("</Relationships>", f'<Relationship Id="{rid}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/font" Target="fonts/font{n}.fntdata"/></Relationships>')
            parts.append(f'<p:{style} r:id="{rid}"/>')
        entries.append(f'<p:embeddedFont><p:font typeface="{fam}"/>{"".join(parts)}</p:embeddedFont>')
    if "fntdata" not in ct:
        ct = ct.replace("<Default ", '<Default Extension="fntdata" ContentType="application/x-fontdata"/><Default ', 1)
    lst = "<p:embeddedFontLst>" + "".join(entries) + "</p:embeddedFontLst>"
    pres = _re.sub(r"(<p:notesSz[^>]*/>)", lambda m: m.group(1) + lst, pres, count=1)
    pres = pres.replace("<p:presentation ", '<p:presentation embedTrueTypeFonts="1" ', 1).replace(' type="screen4x3"', '')
    for item in zin.infolist():
        if item.filename == "ppt/_rels/presentation.xml.rels":
            zout.writestr(item, rels)
        elif item.filename == "ppt/presentation.xml":
            zout.writestr(item, pres)
        elif item.filename == "[Content_Types].xml":
            zout.writestr(item, ct)
        else:
            zout.writestr(item, zin.read(item.filename))
    zout.close(); zin.close()
    shutil.move(tmp, path)
    print("fonts embedded:", n)


embed_fonts(DEST, Path(__file__).parent / "fonts")
